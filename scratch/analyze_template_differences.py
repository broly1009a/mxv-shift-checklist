import os
import openpyxl

ccp_dir = "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Thong ke ccp/output"
ms_lot_dir = "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Thong ke so lot giao dich"
ms_val_dir = "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Thong ke gia tri giao dich"

pairs = [
    ("Thong ke so lot giao dich 2026.xlsx", ccp_dir, "Thong ke so lot giao dich 2026 2.xlsx", ms_lot_dir),
    ("Thong ke so lot giao dich LME 2026.xlsx", ccp_dir, "Thong ke so lot giao dich LME 2026.xlsx", ms_lot_dir),
    ("Thong ke so lot giao dich Options 2026.xlsx", ccp_dir, "Thong ke so lot giao dich Options 2026.xlsx", ms_lot_dir),
    ("Thong ke so lot giao dich Spread 2026.xlsx", ccp_dir, "Thong ke so lot giao dich Spread 2026.xlsx", ms_lot_dir),
    ("Thong ke gia tri giao dich 2026.xlsx", ccp_dir, "Thong ke gia tri giao dich 2026 1.xlsx", ms_val_dir),
    ("Thong ke gia tri giao dich LME 2026.xlsx", ccp_dir, "Thong ke gia tri giao dich LME 2026.xlsx", ms_val_dir),
    ("Thong ke gia tri giao dich Options 2026.xlsx", ccp_dir, "Thong ke gia tri giao dich Options 2026.xlsx", ms_val_dir),
    ("Thong ke gia tri giao dich Spread 2026.xlsx", ccp_dir, "Thong ke gia tri giao dich Spread 2026.xlsx", ms_val_dir),
]

for ccp_f, ccp_d, ms_f, ms_d in pairs:
    ccp_p = os.path.join(ccp_d, ccp_f)
    ms_p = os.path.join(ms_d, ms_f)
    print("=" * 80)
    print(f"FILE: {ccp_f}")
    
    wb_ccp = openpyxl.load_workbook(ccp_p, data_only=True)
    wb_ms = openpyxl.load_workbook(ms_p, data_only=True)
    
    print(f"  Sheets CCP : {wb_ccp.sheetnames}")
    print(f"  Sheets MS  : {wb_ms.sheetnames}")
    
    # Tìm sheet tháng đại diện trong MS
    m_sheets = [s for s in wb_ms.sheetnames if s.startswith('T') and '2026' in s]
    ms_sample_sheet = m_sheets[-1] if m_sheets else wb_ms.sheetnames[0]
    
    # Tìm sheet tháng trong CCP
    ccp_m_sheets = [s for s in wb_ccp.sheetnames if s.startswith('T') and '2026' in s]
    ccp_sample_sheet = ccp_m_sheets[0] if ccp_m_sheets else wb_ccp.sheetnames[0]
    
    ws_ccp = wb_ccp[ccp_sample_sheet]
    ws_ms = wb_ms[ms_sample_sheet]
    
    print(f"  So sánh Sheet Tháng: CCP '{ccp_sample_sheet}' vs MS '{ms_sample_sheet}'")
    print(f"    CCP: max_row={ws_ccp.max_row}, max_col={ws_ccp.max_column}")
    print(f"    MS : max_row={ws_ms.max_row}, max_col={ws_ms.max_column}")
    
    # Lấy row 4 header
    r4_ccp = [str(ws_ccp.cell(4, c).value or "").strip().replace('\\n', ' ') for c in range(1, min(ws_ccp.max_column+1, 20))]
    r4_ms  = [str(ws_ms.cell(4, c).value or "").strip().replace('\\n', ' ') for c in range(1, min(ws_ms.max_column+1, 20))]
    print(f"    Header R4 CCP (10 cols đầu): {r4_ccp[:10]}")
    print(f"    Header R4 MS  (10 cols đầu): {r4_ms[:10]}")
    
    wb_ccp.close()
    wb_ms.close()
