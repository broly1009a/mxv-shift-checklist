import sys
import paramiko

sys.stdout.reconfigure(encoding='utf-8')

ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect('10.0.0.26', username='mxvadmin', password='MxV!,#2o26', timeout=10)

py_script = '''
import openpyxl

# 1. MS TTM
path_ms_ttm = "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/TradingCheck/Futures/2026/T10.2026/20.10/TTM.xlsx"
wb_ms_ttm = openpyxl.load_workbook(path_ms_ttm, read_only=True, data_only=True)
sheet_ms_ttm = wb_ms_ttm.active
rows_ms_ttm = list(sheet_ms_ttm.iter_rows(values_only=True))

ms_ttm_usd = 0
ms_ttm_vnd = 0
for r in rows_ms_ttm[1:]:
    acc = str(r[7] or "").strip()
    if acc.upper().endswith("A"):
        continue
    ms_ttm_usd += float(r[17] or 0)
    ms_ttm_vnd += float(r[18] or 0)

# 2. CCP TTM
path_ccp_ttm = "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Backup CCP/Futures/2026/T10.2026/20.10/TTM CCP.xlsx"
wb_ccp_ttm = openpyxl.load_workbook(path_ccp_ttm, read_only=True, data_only=True)
sheet_ccp_ttm = wb_ccp_ttm.active
rows_ccp_ttm = list(sheet_ccp_ttm.iter_rows(values_only=True))

ccp_ttm_usd = 0
ccp_ttm_vnd = 0
for r in rows_ccp_ttm[1:]:
    ccp_ttm_usd += float(r[13] or 0)
    ccp_ttm_vnd += float(r[14] or 0)

print("=== TỔNG LÃI LỖ DỰ KIẾN (TTM) ===")
print(f"MS TTM (USD - Col 17):  {ms_ttm_usd:>20,f}")
print(f"MS TTM (VND - Col 18):  {ms_ttm_vnd:>20,f}")
print(f"CCP TTM (USD - Col 13): {ccp_ttm_usd:>20,f}")
print(f"CCP TTM (VND - Col 14): {ccp_ttm_vnd:>20,f}")
print("---")
print(f"HIỆN TẠI HỆ THỐNG ĐANG LẤY (USD): {ms_ttm_usd + ccp_ttm_usd:>20,f}")
print(f"ĐÚNG BẢN CHẤT PHẢI LÀ (VND):       {ms_ttm_vnd + ccp_ttm_vnd:>20,f}")

# 3. MS TTTT
path_ms_tttt = "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/TradingCheck/Futures/2026/T10.2026/20.10/TTTT.xlsx"
wb_ms_tttt = openpyxl.load_workbook(path_ms_tttt, read_only=True, data_only=True)
sheet_ms_tttt = wb_ms_tttt.active
rows_ms_tttt = list(sheet_ms_tttt.iter_rows(values_only=True))

ms_tttt_usd = 0
ms_tttt_vnd = 0
for r in rows_ms_tttt[1:]:
    ms_tttt_usd += float(r[19] or 0)
    ms_tttt_vnd += float(r[20] or 0)

# 4. CCP TTTT
path_ccp_tttt = "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Backup CCP/Futures/2026/T10.2026/20.10/TTTT.xlsx"
wb_ccp_tttt = openpyxl.load_workbook(path_ccp_tttt, read_only=True, data_only=True)
sheet_ccp_tttt = wb_ccp_tttt.active
rows_ccp_tttt = list(sheet_ccp_tttt.iter_rows(values_only=True))

ccp_tttt_usd = 0
ccp_tttt_vnd = 0
for r in rows_ccp_tttt[1:]:
    ccp_tttt_usd += float(r[0] or 0)
    ccp_tttt_vnd += float(r[1] or 0)

print("\\n=== LÃI LỖ THỰC TẾ (TTTT) ===")
print(f"MS TTTT (USD - Col 19):  {ms_tttt_usd:>20,f}")
print(f"MS TTTT (VND - Col 20):  {ms_tttt_vnd:>20,f}")
print(f"CCP TTTT (USD - Col 0):  {ccp_tttt_usd:>20,f}")
print(f"CCP TTTT (VND - Col 1):  {ccp_tttt_vnd:>20,f}")
print("---")
print(f"HIỆN TẠI HỆ THỐNG ĐANG LẤY (USD): {ms_tttt_usd + ccp_tttt_usd:>20,f}")
print(f"ĐÚNG BẢN CHẤT PHẢI LÀ (VND):       {ms_tttt_vnd + ccp_tttt_vnd:>20,f}")
'''

stdin, stdout, stderr = ssh.exec_command(f"python3 -c '{py_script}'")
print(stdout.read().decode('utf-8', errors='replace'))
ssh.close()
