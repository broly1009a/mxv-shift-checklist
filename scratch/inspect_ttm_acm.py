import paramiko

ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect('10.0.0.26', port=22, username='mxvadmin', password='MxV!,#2o26', timeout=15)

stdin, stdout, stderr = ssh.exec_command('find /mnt/qlgd-it/TradingCheck/Futures/2026 -name "TTM*.xlsx" | head -n 5')
files = [l.strip() for l in stdout.read().decode().strip().split('\n') if l.strip()]
print('Found TTM files:', files)

if files:
    f = files[0]
    py_cmd = f"""python3 -c "
import openpyxl
wb = openpyxl.load_workbook('{f}', data_only=True)
sheet = wb.active
headers = [c for c in next(sheet.iter_rows(values_only=True))]
print('Headers:', headers[:10])
for idx, r in enumerate(sheet.iter_rows(values_only=True)):
    if idx == 0: continue
    acc = str(r[7] if len(r) > 7 else '')
    if acc.endswith('A') or '-A' in acc:
        print('Found ACM row:', acc, r[:5])
    if idx > 100: break
" """
    stdin, stdout, stderr = ssh.exec_command(py_cmd)
    print(stdout.read().decode())
    print('ERR:', stderr.read().decode())

ssh.close()
