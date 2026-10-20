import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

base_dir = r"M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CE\Futures\2026\T10.2026"
api_file = os.path.join(base_dir, "19.10.api", "DSL ACM CE.xlsx")
ui_file = os.path.join(base_dir, "19.10.giaodien", "DSL ACM CE.xlsx")

print("=== DSL IN 19.10.api ===")
wb = openpyxl.load_workbook(api_file)
ws = wb.active
print("Sheet title:", ws.title)
for r in list(ws.iter_rows(values_only=True))[:5]:
    print(" ", r)

print("\n=== DSL IN 19.10.giaodien ===")
wb_ui = openpyxl.load_workbook(ui_file)
ws_ui = wb_ui.active
print("Sheet title:", ws_ui.title)
for r in list(ws_ui.iter_rows(values_only=True))[:5]:
    print(" ", r)
