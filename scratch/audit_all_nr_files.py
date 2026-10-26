import paramiko

ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect('10.0.0.26', username='mxvadmin', password='MxV!,#2o26', timeout=15)

py_script = """
import openpyxl, glob, json

accs = ['003C3424641-A', '003C8668686-A', '045C8918819-A', '085C8315123-A']

# Let's search all NR.xlsx files in /mnt/qlgd-it for October 2026
nr_files = glob.glob('/mnt/qlgd-it/**/T10.2026/**/NR*.xlsx', recursive=True)
print(f'Found {len(nr_files)} NR files:')
for f in nr_files:
    print(' -', f)

results = []
for f in nr_files:
    try:
        wb = openpyxl.load_workbook(f, data_only=True)
        ws = wb.active
        headers = [str(c.value or '').strip() for c in ws[1]]
        acc_idx = -1
        for i, h in enumerate(headers):
            if any(k in h.lower() for k in ['số tkgd', 'mã tkgd', 'investor code', 'tài khoản']):
                acc_idx = i
                break
        if acc_idx == -1:
            continue
        for row in ws.iter_rows(min_row=2, values_only=True):
            acc = str(row[acc_idx] or '').strip()
            if acc in accs:
                row_dict = {headers[i]: row[i] for i in range(min(len(headers), len(row))) if row[i] is not None}
                row_dict['__source_file__'] = f
                results.append(row_dict)
    except Exception as e:
        print(f'Error reading {f}: {e}')

with open('/tmp/nr_audit.json', 'w', encoding='utf-8') as out:
    json.dump(results, out, ensure_ascii=False, indent=2, default=str)
print(f'Total matching transactions found: {len(results)}')
"""

stdin, stdout, stderr = ssh.exec_command(f'python3 -c "{py_script}"')
print(stdout.read().decode('utf-8'))
print('STDERR:\n', stderr.read().decode('utf-8'))

# Read back /tmp/nr_audit.json
sftp = ssh.open_sftp()
with sftp.file('/tmp/nr_audit.json', 'r') as f:
    audit_data = f.read().decode('utf-8')
sftp.close()
ssh.close()

with open('scratch/nr_audit_result.json', 'w', encoding='utf-8') as f:
    f.write(audit_data)
print('Done saving scratch/nr_audit_result.json')
