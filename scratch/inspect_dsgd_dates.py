import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

base_dir = r"M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CE\Futures\2026\T10.2026"
dsgd_api = os.path.join(base_dir, "19.10.api", "DSGD ACM CE.xlsx")
dsgd_ui = os.path.join(base_dir, "19.10.giaodien", "DSGD ACM CE.xlsx")

print("=== DSGD API ===")
wb = openpyxl.load_workbook(dsgd_api)
ws = wb.active
for r in list(ws.iter_rows(values_only=True))[:5]:
    print(" ", r[:6])

print("\n=== DSGD GIAODIEN ===")
wb2 = openpyxl.load_workbook(dsgd_ui)
ws2 = wb2.active
for r in list(ws2.iter_rows(values_only=True))[:5]:
    print(" ", r[:6])
