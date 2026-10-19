import paramiko

ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect('10.0.0.26', port=22, username='mxvadmin', password='MxV!,#2o26', timeout=15)

py_cmd = """python3 -c "
import openpyxl
f = '/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/TradingCheck/Futures/2026/T10.2026/19.10/TTM.xlsx'
wb = openpyxl.load_workbook(f, data_only=True)
sheet = wb.active
headers = [str(c or '').strip() for c in next(sheet.iter_rows(values_only=True))]
print('Headers:', headers)

acm_count = 0
non_acm_count = 0
acm_examples = []
non_acm_examples = []

# Find column index for account (usually col H or header 'Mã TKGD')
acc_idx = -1
for i, h in enumerate(headers):
    if 'TKGD' in h or 'Tài khoản' in h or 'Số tài khoản' in h:
        acc_idx = i
        break
if acc_idx == -1: acc_idx = 7 # Column H is index 7

print(f'Account column index: {acc_idx} ({headers[acc_idx] if acc_idx < len(headers) else \"N/A\"})')

for idx, r in enumerate(sheet.iter_rows(values_only=True)):
    if idx == 0: continue
    acc = str(r[acc_idx] if len(r) > acc_idx and r[acc_idx] is not None else '').strip()
    if not acc: continue
    
    # Check if ends with A or has -A
    is_acm = acc.upper().endswith('A') or '-A' in acc.upper()
    if is_acm:
        acm_count += 1
        if len(acm_examples) < 5: acm_examples.append(acc)
    else:
        non_acm_count += 1
        if len(non_acm_examples) < 5: non_acm_examples.append(acc)

print(f'Total rows: {acm_count + non_acm_count}')
print(f'ACM accounts (*A): {acm_count}, examples: {acm_examples}')
print(f'Non-ACM accounts (Futures): {non_acm_count}, examples: {non_acm_examples}')
" """

stdin, stdout, stderr = ssh.exec_command(py_cmd)
import sys
res = stdout.read()
sys.stdout.buffer.write(res)
sys.stdout.buffer.write(b'\n')
err_res = stderr.read()
if err_res:
    sys.stdout.buffer.write(b'ERR:\n')
    sys.stdout.buffer.write(err_res)
    sys.stdout.buffer.write(b'\n')
ssh.close()
