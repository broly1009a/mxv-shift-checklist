import paramiko

ssh = paramiko.SSHClient()
ssh.set_missing_host_key_policy(paramiko.AutoAddPolicy())
ssh.connect('10.0.0.26', username='mxvadmin', password='MxV!,#2o26', timeout=15)

py_script = """
import openpyxl, json

base = '/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Backup CCP/Futures/2026/T10.2026/23.10'
qltkgd_path = f'{base}/QL TT TKGD.xlsx'
eod_path = f'{base}/EOD_23.10.2026.xlsx'
nr_path = f'{base}/NR.xlsx'

mismatched_accs = ['003C3424641-A', '003C8668686-A', '045C8918819-A', '085C8315123-A']
neg_accs = ['003C1577943-A', '046C0002363-A', '068C2600030-A']
all_accs = set(mismatched_accs + neg_accs)

print('=== 1. CHECK QL TT TKGD HEADERS & ROWS ===')
wb_ql = openpyxl.load_workbook(qltkgd_path, data_only=True)
ws_ql = wb_ql.active
headers_ql = [str(c.value or '').strip() for c in ws_ql[1]]
print('Headers in QL TT TKGD:', headers_ql)

acc_idx_ql = -1
for i, h in enumerate(headers_ql):
    if any(k in h.lower() for k in ['mã tkgd', 'mã tài khoản', 'mã tk', 'investor code', 'tai khoan']):
        acc_idx_ql = i
        break
print('Acc col idx:', acc_idx_ql, headers_ql[acc_idx_ql] if acc_idx_ql >= 0 else 'NOT FOUND')

for row in ws_ql.iter_rows(min_row=2, values_only=True):
    acc = str(row[acc_idx_ql] or '').strip() if acc_idx_ql >= 0 else ''
    if acc in all_accs:
        row_dict = {headers_ql[i]: row[i] for i in range(len(headers_ql)) if row[i] is not None}
        print(f'QL ROW for {acc}:', json.dumps(row_dict, ensure_ascii=False, default=str))

print('=== 2. CHECK EOD_23.10.2026 HEADERS & ROWS ===')
wb_eod = openpyxl.load_workbook(eod_path, data_only=True)
ws_eod = wb_eod.active
headers_eod = [str(c.value or '').strip() for c in ws_eod[1]]
print('Headers in EOD:', headers_eod)

acc_idx_eod = -1
for i, h in enumerate(headers_eod):
    if any(k in h.lower() for k in ['investorcode', 'investor code', 'mã tkgd', 'account']):
        acc_idx_eod = i
        break

for row in ws_eod.iter_rows(min_row=2, values_only=True):
    acc = str(row[acc_idx_eod] or '').strip() if acc_idx_eod >= 0 else ''
    if acc in all_accs:
        row_dict = {headers_eod[i]: row[i] for i in range(len(headers_eod)) if row[i] is not None}
        print(f'EOD ROW for {acc}:', json.dumps(row_dict, ensure_ascii=False, default=str))

print('=== 3. CHECK NR.xlsx (NỘP RÚT) ===')
try:
    wb_nr = openpyxl.load_workbook(nr_path, data_only=True)
    ws_nr = wb_nr.active
    headers_nr = [str(c.value or '').strip() for c in ws_nr[1]]
    print('Headers in NR:', headers_nr)
    acc_idx_nr = -1
    for i, h in enumerate(headers_nr):
        if any(k in h.lower() for k in ['mã tkgd', 'mã tk', 'investor code', 'tài khoản', 'account']):
            acc_idx_nr = i
            break
    found_nr = 0
    for row in ws_nr.iter_rows(min_row=2, values_only=True):
        acc = str(row[acc_idx_nr] or '').strip() if acc_idx_nr >= 0 else ''
        if acc in all_accs or any(a in str(row) for a in mismatched_accs):
            found_nr += 1
            print(f'NR ROW for {acc}:', [c for c in row if c is not None])
    print('Total matching rows in NR:', found_nr)
except Exception as e:
    print('NR Error:', e)
"""

# write to remote /tmp/inspect_script.py and run
sftp = ssh.open_sftp()
with sftp.file('/tmp/inspect_script.py', 'w') as f:
    f.write(py_script)
sftp.close()

stdin, stdout, stderr = ssh.exec_command('python3 /tmp/inspect_script.py')
out = stdout.read().decode('utf-8')
err = stderr.read().decode('utf-8')
with open('scratch/remote_inspect_result.txt', 'w', encoding='utf-8') as out_f:
    out_f.write('STDOUT:\n' + out + '\nSTDERR:\n' + err)
print('Done writing to scratch/remote_inspect_result.txt')
ssh.close()

