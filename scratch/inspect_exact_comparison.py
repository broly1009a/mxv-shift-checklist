import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

paths = {
    'giaodien': r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CE\Futures\2026\T10.2026\19.10.giaodien',
    'api': r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CE\Futures\2026\T10.2026\19.10.api'
}

files = ['DSGD ACM CE.xlsx', 'DSL ACM CE.xlsx', 'DSLH ACM CE.xlsx', 'DSLCK ACM CE.xlsx', 'DSLDK ACM CE.xlsx']

for f in files:
    print('='*80)
    print(f'FILE: {f}')
    for k, base in paths.items():
        fp = os.path.join(base, f)
        if not os.path.exists(fp):
            print(f'  [{k}] NOT FOUND: {fp}')
            continue
        sz = os.path.getsize(fp)
        wb = openpyxl.load_workbook(fp, data_only=True)
        ws = wb.active
        max_r = ws.max_row
        max_c = ws.max_column
        
        headers = [str(ws.cell(row=1, column=c).value or '') for c in range(1, max_c + 1)]
        
        # Tìm cột ngày hoặc duyệt các dòng để lấy sample
        sample_rows = []
        date_values = set()
        for r in range(2, min(max_r + 1, 100)): # lấy tối đa 100 dòng đầu để xem ngày
            row_vals = [str(ws.cell(row=r, column=c).value or '') for c in range(1, max_c + 1)]
            if r <= 4:
                sample_rows.append(row_vals)
            # Scan all values for date-like strings
            for v in row_vals:
                if '2026' in v or '/10/' in v or '16/10' in v or '19/10' in v:
                    date_values.add(v)
        
        # Cũng quét 10 dòng cuối cùng
        for r in range(max(2, max_r - 10), max_r + 1):
            row_vals = [str(ws.cell(row=r, column=c).value or '') for c in range(1, max_c + 1)]
            for v in row_vals:
                if '2026' in v or '/10/' in v or '16/10' in v or '19/10' in v:
                    date_values.add(v)

        print(f'  [{k}] Size: {sz:,} bytes | Total Rows (incl header): {max_r} | Cols: {max_c}')
        print(f'       Headers (first 6): {headers[:6]}')
        if sample_rows:
            print(f'       Row 2 (first 6): {sample_rows[0][:6]}')
        if len(sample_rows) > 1:
            print(f'       Row 3 (first 6): {sample_rows[1][:6]}')
        print(f'       Dates detected in file: {sorted(list(date_values))[:10]}')
