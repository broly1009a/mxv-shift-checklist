import os
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

base_dir = r"M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CE\Futures\2026\T10.2026"
api_dir = os.path.join(base_dir, "19.10.api")
ui_dir = os.path.join(base_dir, "19.10.giaodien")

print(f"=== SO SÁNH FILE GIỮA 19.10.api VÀ 19.10.giaodien ===")

for name in ["DSL ACM CE.xlsx", "DSLH ACM CE.xlsx", "DSLCK ACM CE.xlsx", "DSLDK ACM CE.xlsx", "DSGD ACM CE.xlsx"]:
    api_file = os.path.join(api_dir, name)
    ui_file = os.path.join(ui_dir, name)

    print(f"\n--- BÁO CÁO: {name} ---")
    if os.path.exists(api_file):
        size_api = os.path.getsize(api_file)
        try:
            wb_api = openpyxl.load_workbook(api_file, read_only=True)
            ws_api = wb_api.active
            rows_api = ws_api.max_row
            print(f"  [API]      Size: {size_api:,} bytes | Max Row: {rows_api}")
        except Exception as e:
            print(f"  [API]      Size: {size_api:,} bytes | Error reading excel: {e}")
    else:
        print(f"  [API]      File không tồn tại!")

    if os.path.exists(ui_file):
        size_ui = os.path.getsize(ui_file)
        try:
            wb_ui = openpyxl.load_workbook(ui_file, read_only=True)
            ws_ui = wb_ui.active
            rows_ui = ws_ui.max_row
            print(f"  [GIAODIEN] Size: {size_ui:,} bytes | Max Row: {rows_ui}")
        except Exception as e:
            print(f"  [GIAODIEN] Size: {size_ui:,} bytes | Error reading excel: {e}")
    else:
        print(f"  [GIAODIEN] File không tồn tại!")
