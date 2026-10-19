import sys
import paramiko

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

client = paramiko.SSHClient()
client.set_missing_host_key_policy(paramiko.AutoAddPolicy())

try:
    client.connect('10.0.0.26', port=22, username='mxvadmin', password='MxV!,#2o26', timeout=10)

    cmd = """python3 -c "
import openpyxl, os

p = '/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Backup CCP/Futures/2026/T09.2026/16.09/GTT CCP.xlsx'
if os.path.exists(p):
    wb = openpyxl.load_workbook(p, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True, max_row=6))
    print('GTT CCP Sheets:', wb.sheetnames)
    print('Header:', rows[0])
    for i in range(1, len(rows)):
        print(f'Row {i+1}:', rows[i])
else:
    print('GTT CCP not found at path')
" """
    stdin, stdout, stderr = client.exec_command(cmd)
    print(stdout.read().decode('utf-8', errors='replace'))
    err = stderr.read().decode('utf-8', errors='replace')
    if err:
        print("ERR:", err)

finally:
    client.close()
