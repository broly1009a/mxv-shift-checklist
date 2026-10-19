import sys
import paramiko

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')

client = paramiko.SSHClient()
client.set_missing_host_key_policy(paramiko.AutoAddPolicy())

try:
    client.connect('10.0.0.26', port=22, username='mxvadmin', password='MxV!,#2o26', timeout=10)

    inspect_ccp_specific = """
import openpyxl, glob, os

files = [
    "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Backup CCP/Futures/2026/T09.2026/16.09/TTM CCP.xlsx",
    "/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Backup CCP/Futures/2026/T09.2026/14.09/DSGD CCP.xlsx",
]
# find any TTTT CCP
extra = glob.glob("/mnt/qlgd-it/Quanlygiaodich/Tai lieu hoat dong/Backup CCP/Futures/*/*/*TTTT*.xlsx")
if extra:
    files.extend(extra[:2])

for f in files:
    if not os.path.exists(f):
        print(f"NOT FOUND: {f}")
        continue
    print(f"\\n==================================================")
    print(f"FILE: {os.path.basename(f)} (Path: {f})")
    wb = openpyxl.load_workbook(f, data_only=True, read_only=True)
    print(f"Sheets: {wb.sheetnames}")
    for sname in wb.sheetnames[:1]:
        ws = wb[sname]
        rows = list(ws.iter_rows(values_only=True, max_row=5))
        print(f"Sheet '{sname}' Header (Row 1):")
        print(rows[0] if len(rows) > 0 else "EMPTY")
        if len(rows) > 1:
            print("Row 2 (Data 1):", rows[1])
        if len(rows) > 2:
            print("Row 3 (Data 2):", rows[2])
"""
    sftp = client.open_sftp()
    with sftp.file('/tmp/inspect_ccp_ttm.py', 'w') as f:
        f.write(inspect_ccp_specific)
    sftp.close()

    stdin, stdout, stderr = client.exec_command('python3 /tmp/inspect_ccp_ttm.py')
    print(stdout.read().decode('utf-8', errors='replace'))
    err = stderr.read().decode('utf-8', errors='replace')
    if err:
        print("ERR:", err)

finally:
    client.close()
