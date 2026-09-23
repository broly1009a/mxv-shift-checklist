import os
import sys
import io
import json
import openpyxl
from collections import defaultdict

sys.stdout.reconfigure(encoding='utf-8')

dir_path = os.path.dirname(os.path.abspath(__file__))

def load_wb(fname):
    p = os.path.join(dir_path, fname)
    with open(p, 'rb') as f:
        file_bytes = io.BytesIO(f.read())
    return openpyxl.load_workbook(file_bytes, data_only=True)

# 1. Load Commodity Specs
wb_hh = load_wb('HH.csv')
sheet_hh = wb_hh.active
hh_specs = {}
for r in range(2, sheet_hh.max_row + 1):
    ma_hh = str(sheet_hh.cell(r, 2).value or '').strip().upper()
    ten_hh = str(sheet_hh.cell(r, 6).value or '').strip()
    do_cao = float(sheet_hh.cell(r, 8).value or 1)
    tien_te = str(sheet_hh.cell(r, 15).value or 'USD').strip().upper()
    if ma_hh:
        hh_specs[ma_hh] = {
            'maHH': ma_hh,
            'tenHH': ten_hh,
            'doCao': do_cao,
            'tienTe': tien_te
        }

# 2. Extract Commodity from Contract code
def get_ma_hh(ma_hd):
    clean = str(ma_hd or '').strip().upper()
    for known in ['SI5CO', 'CP2CO', 'PL1NY', 'SIV']:
        if clean.startswith(known):
            return known
    if len(clean) >= 5:
        return clean[:5]
    return clean

# 3. Parse DSGD.csv
wb_dsgd = load_wb('DSGD.csv')
sheet_dsgd = wb_dsgd.active

header_dsgd = [sheet_dsgd.cell(1, c).value for c in range(1, sheet_dsgd.max_column + 1)]
print(f"DSGD Header: {header_dsgd}")

dsgd_rows = []
for r in range(2, sheet_dsgd.max_row + 1):
    ma_tk = str(sheet_dsgd.cell(r, 6).value or '').strip()
    ma_hd = str(sheet_dsgd.cell(r, 7).value or '').strip()
    mua_ban = str(sheet_dsgd.cell(r, 8).value or '').strip()
    loai_lenh = str(sheet_dsgd.cell(r, 9).value or '').strip().upper()
    kl_khop = float(sheet_dsgd.cell(r, 11).value or 0)
    gia_khop = float(sheet_dsgd.cell(r, 14).value or 0)
    tvkd = str(sheet_dsgd.cell(r, 24).value or '').strip()
    if len(tvkd) < 3 and tvkd.isdigit():
        tvkd = tvkd.zfill(3)
    
    ma_hh = get_ma_hh(ma_hd)
    spec = hh_specs.get(ma_hh, {'doCao': 1, 'tienTe': 'USD', 'tenHH': ma_hh})
    do_cao = spec['doCao']
    
    # Calculate GTGD at 25920 and 26000
    gtgd_raw_usd = kl_khop * gia_khop * do_cao
    gtgd_25920 = round(gtgd_raw_usd * 25920)
    gtgd_26000 = round(gtgd_raw_usd * 26000)
    
    dsgd_rows.append({
        'row_idx': r,
        'maTKGD': ma_tk,
        'tvkd': tvkd,
        'maHD': ma_hd,
        'maHH': ma_hh,
        'tenHH': spec['tenHH'],
        'muaBan': mua_ban,
        'loaiLenh': loai_lenh,
        'klKhop': kl_khop,
        'giaKhop': gia_khop,
        'doCao': do_cao,
        'gtgd_raw_usd': gtgd_raw_usd,
        'gtgd_25920': gtgd_25920,
        'gtgd_26000': gtgd_26000,
        'isAcm': ma_tk.endswith('-A')
    })

# 4. Parse TTM.csv
wb_ttm = load_wb('TTM.csv')
sheet_ttm = wb_ttm.active
ttm_by_tvkd = defaultdict(lambda: {'klMua': 0, 'klBan': 0, 'laiLoVnd': 0, 'count': 0})

for r in range(2, sheet_ttm.max_row + 1):
    tvkd = str(sheet_ttm.cell(r, 1).value or '').strip()
    if len(tvkd) < 3 and tvkd.isdigit():
        tvkd = tvkd.zfill(3)
    kl_mua = float(sheet_ttm.cell(r, 9).value or 0)
    kl_ban = float(sheet_ttm.cell(r, 10).value or 0)
    lai_lo_vnd = float(sheet_ttm.cell(r, 15).value or 0)
    
    ttm_by_tvkd[tvkd]['klMua'] += kl_mua
    ttm_by_tvkd[tvkd]['klBan'] += kl_ban
    ttm_by_tvkd[tvkd]['laiLoVnd'] += lai_lo_vnd
    ttm_by_tvkd[tvkd]['count'] += 1

# 5. Parse TTTT.csv
wb_tttt = load_wb('TTTT.csv')
sheet_tttt = wb_tttt.active
tttt_by_tvkd = defaultdict(lambda: {'klMua': 0, 'klBan': 0, 'laiLoVnd': 0, 'count': 0})

for r in range(2, sheet_tttt.max_row + 1):
    tvkd = str(sheet_tttt.cell(r, 1).value or '').strip()
    if len(tvkd) < 3 and tvkd.isdigit():
        tvkd = tvkd.zfill(3)
    lai_lo_vnd = float(sheet_tttt.cell(r, 6).value or 0)
    kl_mua = float(sheet_tttt.cell(r, 10).value or 0)
    kl_ban = float(sheet_tttt.cell(r, 11).value or 0)
    
    tttt_by_tvkd[tvkd]['klMua'] += kl_mua
    tttt_by_tvkd[tvkd]['klBan'] += kl_ban
    tttt_by_tvkd[tvkd]['laiLoVnd'] += lai_lo_vnd
    tttt_by_tvkd[tvkd]['count'] += 1

# 6. Aggregate by TVKD
tvkd_stats = defaultdict(lambda: {
    'soLot': 0,
    'gtgd_25920': 0,
    'gtgd_26000': 0,
    'gtgd_raw_usd': 0,
    'orderTypes': set(),
    'byHH': defaultdict(lambda: {'soLot': 0, 'gtgd_25920': 0, 'gtgd_26000': 0}),
    'tradeCount': 0
})

for row in dsgd_rows:
    t = tvkd_stats[row['tvkd']]
    t['tradeCount'] += 1
    t['soLot'] += row['klKhop']
    t['gtgd_25920'] += row['gtgd_25920']
    t['gtgd_26000'] += row['gtgd_26000']
    t['gtgd_raw_usd'] += row['gtgd_raw_usd']
    t['orderTypes'].add(row['loaiLenh'])
    
    hh = t['byHH'][row['maHH']]
    hh['soLot'] += row['klKhop']
    hh['gtgd_25920'] += row['gtgd_25920']
    hh['gtgd_26000'] += row['gtgd_26000']

# Print Detailed Analysis
print(f"\nTotal DSGD trades parsed: {len(dsgd_rows)}")
total_lot = sum(r['klKhop'] for r in dsgd_rows)
total_gtgd_25920 = sum(r['gtgd_25920'] for r in dsgd_rows)
total_gtgd_26000 = sum(r['gtgd_26000'] for r in dsgd_rows)
total_usd = sum(r['gtgd_raw_usd'] for r in dsgd_rows)

print(f"Total Lots: {total_lot}")
print(f"Total USD: {total_usd:,.2f} USD")
print(f"Total GTGD (Rate 25,920): {total_gtgd_25920:,.0f} VND")
print(f"Total GTGD (Rate 26,000): {total_gtgd_26000:,.0f} VND")
print(f"Difference: {total_gtgd_26000 - total_gtgd_25920:,.0f} VND")

# Save full summary json
summary = {
    'totalTrades': len(dsgd_rows),
    'totalLot': total_lot,
    'totalUsd': total_usd,
    'totalGtgd25920': total_gtgd_25920,
    'totalGtgd26000': total_gtgd_26000,
    'diffVnd': total_gtgd_26000 - total_gtgd_25920,
    'tvkd': {}
}

for tvkd in sorted(tvkd_stats.keys()):
    st = tvkd_stats[tvkd]
    ttm = ttm_by_tvkd.get(tvkd, {'klMua': 0, 'klBan': 0, 'laiLoVnd': 0})
    tttt = tttt_by_tvkd.get(tvkd, {'klMua': 0, 'klBan': 0, 'laiLoVnd': 0})
    
    missing_types = [ot for ot in ['MKT', 'LMT', 'STP', 'STL'] if ot not in st['orderTypes']]
    
    summary['tvkd'][tvkd] = {
        'soLot': st['soLot'],
        'gtgd_25920': st['gtgd_25920'],
        'gtgd_26000': st['gtgd_26000'],
        'gtgd_usd': st['gtgd_raw_usd'],
        'tradeCount': st['tradeCount'],
        'orderTypes': list(st['orderTypes']),
        'isFull4Types': len(missing_types) == 0,
        'missingTypes': missing_types,
        'ttm': {
            'klMua': ttm['klMua'],
            'klBan': ttm['klBan'],
            'totalLot': ttm['klMua'] + ttm['klBan'],
            'laiLoVnd': ttm['laiLoVnd']
        },
        'tttt': {
            'klMua': tttt['klMua'],
            'klBan': tttt['klBan'],
            'totalLot': tttt['klMua'] + tttt['klBan'],
            'laiLoVnd': tttt['laiLoVnd']
        },
        'byHH': dict(st['byHH'])
    }

with open(os.path.join(dir_path, 'python_verification_result.json'), 'w', encoding='utf-8') as f:
    json.dump(summary, f, ensure_ascii=False, indent=2)

print("\nSaved result to python_verification_result.json")
