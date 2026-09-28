import openpyxl
import sys

sys.stdout.reconfigure(encoding='utf-8')

p1 = r'c:\Users\hiepth\OneDrive - MERCANTILE EXCHANGE OF VIETNAM\Documents\Github\mxv-cqg-download-investigation\backend\src\modules\ccp-statistics\outputACM24.09\Thong ke gia tri giao dich ACM 2026 (Theo CCP).xlsx'
p2 = r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Thong ke ccp\output\Thong ke gia tri giao dich ACM 2026.xlsx'

wb1 = openpyxl.load_workbook(p1, data_only=True)
ws1 = wb1['T09.2026']

wb2 = openpyxl.load_workbook(p2, data_only=True)
ws2 = wb2['T09.2026']

print(f"{'Row':<4} | {'Date':<10} | {'File 1 (User / OutputACM)':<65} | {'File 2 (Output He Thong)':<65}")
print(f"{'':<4} | {'':<10} | {'SI5CO':>14} {'PL1NY':>14} {'CP2CO':>14} {'Tong':>18} | {'SI5CO':>14} {'PL1NY':>14} {'CP2CO':>14} {'Tong':>18}")
print('-' * 155)

for r in range(15, 26):
    d1 = str(ws1.cell(r, 1).value or '')[:10]
    v1_b = ws1.cell(r, 2).value or 0
    v1_c = ws1.cell(r, 3).value or 0
    v1_d = ws1.cell(r, 4).value or 0
    v1_e = ws1.cell(r, 5).value or 0

    v2_b = ws2.cell(r, 2).value or 0
    v2_c = ws2.cell(r, 3).value or 0
    v2_d = ws2.cell(r, 4).value or 0
    v2_e = ws2.cell(r, 5).value or 0

    s1 = f"{v1_b:>14,} {v1_c:>14,} {v1_d:>14,} {v1_e:>18,}"
    s2 = f"{v2_b:>14,} {v2_c:>14,} {v2_d:>14,} {v2_e:>18,}"

    flag = "  <-- MATCH" if s1 == s2 else "  <-- DIFF!"
    print(f"{r:<4} | {d1:<10} | {s1} | {s2} {flag}")
