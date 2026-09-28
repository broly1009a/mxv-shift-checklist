import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

# Check 21.09 MS
ms_21 = r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup MS\Futures\2026\T09.2026\21.09\DSGD.xlsx'
if os.path.exists(ms_21):
    wb = openpyxl.load_workbook(ms_21, data_only=True)
    ws = wb.active
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    h_idx = {h: i+1 for i, h in enumerate(headers)}
    
    si = 0; pl = 0; cp = 0
    for r in range(2, ws.max_row + 1):
        mahd = str(ws.cell(r, h_idx['Mã HĐ']).value or '')
        kl = float(ws.cell(r, h_idx['KL giao dịch']).value or 0)
        gia = float(ws.cell(r, h_idx['Giá khớp']).value or 0)
        if 'SI5CO' in mahd: si += kl * gia * 100 * 26000
        elif 'PL1NY' in mahd: pl += kl * gia * 5 * 26000
        elif 'CP2CO' in mahd: cp += kl * gia * 1000 * 26000
    print(f"21.09 MS DSGD: SI={si:,.0f}, PL={pl:,.0f}, CP={cp:,.0f}")

# Check 21.09 ACM (Straits.csv)
# Check if there is another file
