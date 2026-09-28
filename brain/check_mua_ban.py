import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

backup_dir = r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CCP\Futures\2026\T09.2026\24.09'
dsgd_path = os.path.join(backup_dir, 'DSGD CCP.xlsx')

wb = openpyxl.load_workbook(dsgd_path, data_only=True)
ws = wb.active

headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
h_idx = {h: i+1 for i, h in enumerate(headers)}

rows = []
for r in range(2, ws.max_row + 1):
    row_data = {h: ws.cell(r, c).value for h, c in h_idx.items()}
    rows.append(row_data)

# Let's inspect by Mua / Bán
for hh in ['SI5CO', 'PL1NY', 'CP2CO']:
    hh_rows = [r for r in rows if hh in str(r.get('Mã HĐ'))]
    print(f"\n=== {hh} ({len(hh_rows)} rows) ===")
    
    # Check by Mua/Bán
    mua_rows = [r for r in hh_rows if str(r.get('Mua/Bán')).strip() == 'Mua']
    ban_rows = [r for r in hh_rows if str(r.get('Mua/Bán')).strip() == 'Bán']
    
    docao = 100 if hh == 'SI5CO' else (5 if hh == 'PL1NY' else 1000)
    
    mua_val = sum(float(r.get('KL khớp') or 0) * float(r.get('Giá khớp trung bình') or 0) for r in mua_rows) * docao
    ban_val = sum(float(r.get('KL khớp') or 0) * float(r.get('Giá khớp trung bình') or 0) for r in ban_rows) * docao
    all_val = sum(float(r.get('KL khớp') or 0) * float(r.get('Giá khớp trung bình') or 0) for r in hh_rows) * docao
    
    print(f"  All: {all_val:,.2f} USD -> {all_val * 26000:,.0f} VND (KL={sum(float(r.get('KL khớp') or 0) for r in hh_rows)})")
    print(f"  Mua: {mua_val:,.2f} USD -> {mua_val * 26000:,.0f} VND (KL={sum(float(r.get('KL khớp') or 0) for r in mua_rows)})")
    print(f"  Ban: {ban_val:,.2f} USD -> {ban_val * 26000:,.0f} VND (KL={sum(float(r.get('KL khớp') or 0) for r in ban_rows)})")
