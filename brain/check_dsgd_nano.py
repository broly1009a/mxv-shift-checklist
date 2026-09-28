import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

backup_dir = r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CCP\Futures\2026\T09.2026\24.09'

# 1. Inspect DSGD CCP.xlsx
dsgd_path = os.path.join(backup_dir, 'DSGD CCP.xlsx')
wb = openpyxl.load_workbook(dsgd_path, data_only=True)
ws = wb.active
print(f"=== DSGD CCP.xlsx (rows: {ws.max_row}) ===")
headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
print("Headers:", headers)

rows_nano = []
for r in range(2, ws.max_row + 1):
    row_vals = {headers[c-1]: ws.cell(r, c).value for c in range(1, ws.max_column + 1)}
    # check if nano
    tkgd = str(row_vals.get('Mã TKGD') or row_vals.get('Tài khoản') or '')
    mahd = str(row_vals.get('Mã HĐ') or row_vals.get('Hợp đồng') or '')
    if any(k in mahd for k in ['SI5CO', 'PL1NY', 'CP2CO']) or tkgd.endswith('-A') or tkgd.endswith('A'):
        rows_nano.append(row_vals)

print(f"Total nano rows in DSGD CCP: {len(rows_nano)}")
for r in rows_nano:
    print(r)
