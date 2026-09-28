import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

backup_dir = r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CCP\Futures\2026\T09.2026\24.09'
dsgd_path = os.path.join(backup_dir, 'DSGD CCP.xlsx')

wb = openpyxl.load_workbook(dsgd_path, data_only=True)
ws = wb.active

headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
h_idx = {h: i+1 for i, h in enumerate(headers)}

rows = []
for r in range(2, ws.max_row + 1):
    row_data = {h: ws.cell(r, c).value for h, c in h_idx.items()}
    rows.append(row_data)

print(f"Total rows in DSGD CCP.xlsx: {len(rows)}")

# Check commodities
from collections import defaultdict
by_commodity = defaultdict(list)
for r in rows:
    ma_hd = str(r.get('Mã HĐ') or '')
    san = str(r.get('Sàn giao dịch') or '')
    tkgd = str(r.get('Mã TKGD') or '')
    
    # get prefix
    ma_hh = ma_hd[:5] if len(ma_hd) >= 5 else ma_hd
    by_commodity[ma_hh].append(r)

print("\nRows by commodity prefix:")
for k, v in by_commodity.items():
    print(f"  {k}: {len(v)} rows, total KL={sum(float(x.get('KL khớp') or 0) for x in v)}")
    for x in v[:2]:
        print(f"     Example: TK={x.get('Mã TKGD')}, HD={x.get('Mã HĐ')}, KL={x.get('KL khớp')}, Gia={x.get('Giá khớp trung bình')}, San={x.get('Sàn giao dịch')}")

# Check Tỷ giá in backup_dir
for f in os.listdir(backup_dir):
    if 'TG' in f or 'TYGIA' in f or 'EXCHANGE' in f.upper() or 'RATE' in f.upper() or 'CURRENCY' in f.upper():
        print(f"Found exchange rate file: {f}")
