import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

search_dir = r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Thong ke ccp'
targets = [8516274000, 468767000, 21695076000, 30680117000]

print("Searching in:", search_dir)
for root, dirs, files in os.walk(search_dir):
    for f in files:
        if not f.endswith('.xlsx'): continue
        fpath = os.path.join(root, f)
        try:
            wb = openpyxl.load_workbook(fpath, data_only=True)
            for s in wb.sheetnames:
                ws = wb[s]
                for r in range(1, ws.max_row + 1):
                    # Check cell values
                    for c in range(1, ws.max_column + 1):
                        val = ws.cell(r, c).value
                        if val in targets:
                            print(f"FOUND MATCH in {os.path.relpath(fpath, search_dir)} -> Sheet '{s}' -> Cell({r},{c}) = {val}")
        except Exception as e:
            pass
print("Done search in Thong ke ccp.")
