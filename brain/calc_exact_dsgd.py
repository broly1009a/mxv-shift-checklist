import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

backup_dir = r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CCP\Futures\2026\T09.2026\24.09'
dsgd_path = os.path.join(backup_dir, 'DSGD CCP.xlsx')

wb = openpyxl.load_workbook(dsgd_path, data_only=True)
ws = wb.active

headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
h_idx = {h: i+1 for i, h in enumerate(headers)}

rows = []
for r in range(2, ws.max_row + 1):
    row_data = {h: ws.cell(r, c).value for h, c in h_idx.items()}
    rows.append(row_data)

# Let's inspect each commodity's rows
# SI5CO:
si_rows = [r for r in rows if 'SI5CO' in str(r.get('Mã HĐ'))]
pl_rows = [r for r in rows if 'PL1NY' in str(r.get('Mã HĐ'))]
cp_rows = [r for r in rows if 'CP2CO' in str(r.get('Mã HĐ'))]

print(f"SI5CO: {len(si_rows)} rows")
print(f"PL1NY: {len(pl_rows)} rows")
print(f"CP2CO: {len(cp_rows)} rows")

# Check CP2CO rows:
print("\n--- CP2CO rows (2 rows) ---")
for r in cp_rows:
    print(r.get('Mã TKGD'), r.get('Mã HĐ'), "KL:", r.get('KL khớp'), "Gia:", r.get('Giá khớp trung bình'))

# Check PL1NY rows:
print("\n--- PL1NY rows (4 rows) ---")
for r in pl_rows:
    print(r.get('Mã TKGD'), r.get('Mã HĐ'), "KL:", r.get('KL khớp'), "Gia:", r.get('Giá khớp trung bình'))

# Check SI5CO summary:
print("\n--- SI5CO rows summary ---")
sum_kl_gia = sum(float(r.get('KL khớp') or 0) * float(r.get('Giá khớp trung bình') or 0) for r in si_rows)
sum_kl = sum(float(r.get('KL khớp') or 0) for r in si_rows)
print(f"SI5CO sum(KL * Gia) = {sum_kl_gia:.4f}, sum(KL) = {sum_kl}")

# What exchange rate is used in system?
# Let's check what exchange rate file exists or what USD rate is used in CCP
# In system: tyGia(USD) might be ~26,000 or similar
for usd_rate in [25920, 26000, 25900, 25000, 25800]:
    # SI5CO: doCao = ?
    # PL1NY: doCao = ?
    # CP2CO: doCao = ?
    pass
