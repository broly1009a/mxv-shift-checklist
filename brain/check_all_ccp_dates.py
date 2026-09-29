import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

dates = ['21.09', '22.09', '23.09', '24.09', '25.09']

for d in dates:
    print(f"\n==================== DATE {d} ====================")
    # Check CCP DSGD
    ccp_dir = rf'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CCP\Futures\2026\T09.2026\{d}'
    dsgd_ccp = None
    for fname in ['DSGD CCP.xlsx', f'DSGD_{d}.2026.xlsx', 'DSGD.xlsx']:
        p = os.path.join(ccp_dir, fname)
        if os.path.exists(p):
            dsgd_ccp = p
            break
    
    if dsgd_ccp:
        wb = openpyxl.load_workbook(dsgd_ccp, data_only=True)
        ws = wb.active
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        h_idx = {h: i+1 for i, h in enumerate(headers)}
        
        si_kl = 0; si_val = 0
        pl_kl = 0; pl_val = 0
        cp_kl = 0; cp_val = 0
        
        for r in range(2, ws.max_row + 1):
            mahd = str(ws.cell(r, h_idx.get('Mã HĐ') or h_idx.get('Hợp đồng') or 1).value or '')
            kl = float(ws.cell(r, h_idx.get('KL khớp') or h_idx.get('Khối lượng') or 1).value or 0)
            gia = float(ws.cell(r, h_idx.get('Giá khớp trung bình') or h_idx.get('Giá khớp') or 1).value or 0)
            
            if 'SI5CO' in mahd:
                si_kl += kl; si_val += kl * gia * 100 * 26000
            elif 'PL1NY' in mahd:
                pl_kl += kl; pl_val += kl * gia * 5 * 26000
            elif 'CP2CO' in mahd:
                cp_kl += kl; cp_val += kl * gia * 1000 * 26000
                
        print(f"CCP DSGD ({os.path.basename(dsgd_ccp)}):")
        print(f"  SI5CO: {si_kl:4.0f} lots -> {si_val:>15,.0f} VND")
        print(f"  PL1NY: {pl_kl:4.0f} lots -> {pl_val:>15,.0f} VND")
        print(f"  CP2CO: {cp_kl:4.0f} lots -> {cp_val:>15,.0f} VND")
        print(f"  TOTAL: {si_kl+pl_kl+cp_kl:4.0f} lots -> {si_val+pl_val+cp_val:>15,.0f} VND")
    else:
        print("No CCP DSGD found")
