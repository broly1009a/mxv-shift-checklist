import openpyxl
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

backup_dir = r'M:\Tailieuchung\QLGD-IT\Quanlygiaodich\Tai lieu hoat dong\Backup CCP\Futures\2026\T09.2026\24.09'

targets = [
    8516274000, 468767000, 21695076000, 30680117000,
    327549, 18029.5, 834426,
    13858650000, 917332000, 352170000, 15128152000
]

for f in sorted(os.listdir(backup_dir)):
    if not f.endswith('.xlsx'): continue
    fpath = os.path.join(backup_dir, f)
    try:
        wb = openpyxl.load_workbook(fpath, data_only=True)
        for s in wb.sheetnames:
            ws = wb[s]
            for r in range(1, min(ws.max_row + 1, 500)):
                for c in range(1, min(ws.max_column + 1, 50)):
                    val = ws.cell(r, c).value
                    if val is None: continue
                    # check if number
                    try:
                        v_num = float(val)
                        for t in targets:
                            if abs(v_num - t) < 1:
                                print(f"MATCH in {f} -> Sheet '{s}' -> Cell({r},{c}) = {val} (target {t})")
                    except:
                        pass
    except Exception as e:
        # print(f"Error {f}: {e}")
        pass
